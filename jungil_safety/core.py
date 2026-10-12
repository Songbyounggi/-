"""정일건설 안전관리 도우미 - 데이터/로직 모듈 (GUI 없음).

HPMS '안전·보건 부적합 보고서(현장)' 와 'TBM 보고서' 화면 항목을 그대로 따라
협력업체(정일건설) 측에서 자체 관리할 수 있도록 SQLite 에 저장한다.
"""
import datetime as dt
import os
import re
import sqlite3
import sys

COMPANY = "정일건설 주식회사"
DEFAULT_SITE_CODE = "I30F"
DEFAULT_SITE_NAME = "태안기업도시 기반시설공사[4차]"

# (컬럼키, 화면표시명, 기본폭)
NCR_FIELDS = [
    ("report_no", "보고서 번호", 190),
    ("dept", "발행부서", 70),
    ("source", "출처", 60),
    ("author", "작성자", 70),
    ("issue_date", "발행일", 90),
    ("due_date", "조치요구일", 90),
    ("done_date", "조치완료", 90),
    ("receiver", "수신처", 130),
    ("issue_type", "발행구분", 110),
    ("title", "제목", 260),
    ("result", "결과확인", 70),
    ("location", "위치/작업", 120),
    ("action", "조치내용", 200),
    ("memo", "비고", 120),
]
TBM_FIELDS = [
    ("site_code", "현장코드", 70),
    ("ptw_no", "PTW 관리번호", 140),
    ("company", "협력업체명", 130),
    ("reg_date", "등록일", 90),
    ("exec_date", "실시일", 90),
    ("kind", "실시구분", 110),
    ("work", "작업내용", 220),
    ("leader", "작업반장", 80),
    ("attendees", "참석인원", 70),
    ("hazards", "위험요인", 200),
    ("measures", "안전대책", 200),
    ("memo", "비고", 120),
]

DEPTS = ["공무", "토목", "안전", "공사", "품질", "환경"]
SOURCES = ["일반", "본사", "발주처", "감리"]
ISSUE_TYPES = ["현장 자체 점검", "일일통합점검", "합동점검", "본사 점검", "감리 지적"]
RESULTS = ["완료", "미완료"]
TBM_KINDS = ["필수(최초TBM)", "일반TBM", "수시TBM"]

DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def app_dir():
    """실행파일(또는 스크립트) 옆 폴더 - USB 등에 통째로 들고 다닐 수 있게."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def today():
    return dt.date.today().isoformat()


def parse_date(value):
    """'2026-09-30', '2026.09.30', '20260930', datetime 등 -> 'YYYY-MM-DD' (빈값은 '')."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    s = str(value).strip()
    if not s:
        return ""
    s = s.split(" ")[0].replace(".", "-").replace("/", "-")
    if re.fullmatch(r"\d{8}", s):
        s = f"{s[:4]}-{s[4:6]}-{s[6:]}"
    try:
        return dt.date.fromisoformat(
            "-".join(p.zfill(2) for p in s.split("-"))
        ).isoformat()
    except ValueError:
        raise ValueError(f"날짜 형식 오류: {value!r} (예: 2026-09-30)")


def ncr_status(row, ref=None):
    """완료 / 진행중 / 기한초과."""
    ref = ref or today()
    if row["done_date"]:
        return "완료"
    if row["due_date"] and row["due_date"] < ref:
        return "기한초과"
    return "진행중"


def on_time_rate(rows, ref=None):
    """기한 내 조치율(%) - 조치기한이 지났거나 이미 조치된 건 중 기한 내 완료 비율.

    대상이 없으면 None.
    """
    ref = ref or today()
    target = [r for r in rows if r["done_date"] or (r["due_date"] and r["due_date"] < ref)]
    if not target:
        return None
    ok = [
        r for r in target
        if r["done_date"] and (not r["due_date"] or r["done_date"] <= r["due_date"])
    ]
    return round(len(ok) * 100.0 / len(target), 2)


def _seq(report_no):
    tail = report_no.rsplit("-", 1)[-1]
    return int(tail) if tail.isdigit() else 0


class Store:
    def __init__(self, path=None):
        self.path = path or os.path.join(app_dir(), "jungil_safety.db")
        self.db = sqlite3.connect(self.path)
        self.db.row_factory = sqlite3.Row
        self._init()

    # ---------------------------------------------------------------- schema
    def _init(self):
        ncr_cols = ", ".join(f"{k} TEXT DEFAULT ''" for k, _, _ in NCR_FIELDS)
        tbm_cols = ", ".join(f"{k} TEXT DEFAULT ''" for k, _, _ in TBM_FIELDS)
        self.db.executescript(
            f"""
            CREATE TABLE IF NOT EXISTS ncr (id INTEGER PRIMARY KEY AUTOINCREMENT, {ncr_cols});
            CREATE UNIQUE INDEX IF NOT EXISTS ncr_no ON ncr(report_no);
            CREATE TABLE IF NOT EXISTS tbm (id INTEGER PRIMARY KEY AUTOINCREMENT, {tbm_cols});
            CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT);
            """
        )
        for k, v in (
            ("site_code", DEFAULT_SITE_CODE),
            ("site_name", DEFAULT_SITE_NAME),
            ("company", COMPANY),
        ):
            self.db.execute("INSERT OR IGNORE INTO settings VALUES (?, ?)", (k, v))
        self.db.commit()

    def setting(self, key):
        r = self.db.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
        return r["value"] if r else ""

    def set_setting(self, key, value):
        self.db.execute("INSERT OR REPLACE INTO settings VALUES (?, ?)", (key, value))
        self.db.commit()

    # ------------------------------------------------------------------- NCR
    def next_report_no(self, issue_date):
        """I30F-HSE-NCR-202609-11 형식으로 해당 월 다음 번호."""
        ym = parse_date(issue_date)[:7].replace("-", "")
        prefix = f"{self.setting('site_code')}-HSE-NCR-{ym}-"
        n = 0
        for r in self.db.execute("SELECT report_no FROM ncr WHERE report_no LIKE ?", (prefix + "%",)):
            tail = r["report_no"][len(prefix):]
            if tail.isdigit():
                n = max(n, int(tail))
        return f"{prefix}{n + 1}"

    def _clean(self, fields, data, date_keys):
        out = {}
        for k, _, _ in fields:
            v = data.get(k, "")
            v = "" if v is None else str(v).strip()
            if k in date_keys:
                v = parse_date(v)
            out[k] = v
        return out

    def save_ncr(self, data, rid=None):
        d = self._clean(NCR_FIELDS, data, ("issue_date", "due_date", "done_date"))
        if not d["title"]:
            raise ValueError("제목을 입력하세요.")
        if not d["issue_date"]:
            d["issue_date"] = today()
        if not d["report_no"]:
            d["report_no"] = self.next_report_no(d["issue_date"])
        if d["due_date"] and d["due_date"] < d["issue_date"]:
            raise ValueError("조치요구일이 발행일보다 빠릅니다.")
        if not d["receiver"]:
            d["receiver"] = self.setting("company")
        if d["done_date"] and not d["result"]:
            d["result"] = "완료"
        return self._upsert("ncr", NCR_FIELDS, d, rid)

    def list_ncr(self, start="", end="", status="전체", receiver="", keyword=""):
        sql, args = "SELECT * FROM ncr WHERE 1=1", []
        if start:
            sql += " AND issue_date >= ?"
            args.append(parse_date(start))
        if end:
            sql += " AND issue_date <= ?"
            args.append(parse_date(end))
        if receiver:
            sql += " AND receiver LIKE ?"
            args.append(f"%{receiver}%")
        if keyword:
            sql += " AND (title LIKE ? OR report_no LIKE ? OR author LIKE ? OR action LIKE ?)"
            args += [f"%{keyword}%"] * 4
        sql += " ORDER BY issue_date DESC, id DESC"
        rows = [dict(r) for r in self.db.execute(sql, args)]
        # 같은 발행일은 보고서 순번 큰 것부터 (HPMS 목록과 같은 순서)
        rows.sort(key=lambda r: (r["issue_date"], _seq(r["report_no"]), r["id"]), reverse=True)
        for r in rows:
            r["status"] = ncr_status(r)
        if status and status != "전체":
            rows = [r for r in rows if r["status"] == status]
        return rows

    # ------------------------------------------------------------------- TBM
    def save_tbm(self, data, rid=None):
        d = self._clean(TBM_FIELDS, data, ("reg_date", "exec_date"))
        if not d["exec_date"]:
            raise ValueError("실시일을 입력하세요.")
        if not d["reg_date"]:
            d["reg_date"] = d["exec_date"]
        if not d["site_code"]:
            d["site_code"] = self.setting("site_code")
        if not d["company"]:
            d["company"] = self.setting("company")
        if d["attendees"] and not d["attendees"].isdigit():
            raise ValueError("참석인원은 숫자로 입력하세요.")
        return self._upsert("tbm", TBM_FIELDS, d, rid)

    def list_tbm(self, start="", end="", kind="전체", keyword=""):
        sql, args = "SELECT * FROM tbm WHERE 1=1", []
        if start:
            sql += " AND exec_date >= ?"
            args.append(parse_date(start))
        if end:
            sql += " AND exec_date <= ?"
            args.append(parse_date(end))
        if kind and kind != "전체":
            sql += " AND kind = ?"
            args.append(kind)
        if keyword:
            sql += " AND (ptw_no LIKE ? OR work LIKE ? OR company LIKE ? OR leader LIKE ?)"
            args += [f"%{keyword}%"] * 4
        sql += " ORDER BY exec_date DESC, id DESC"
        return [dict(r) for r in self.db.execute(sql, args)]

    def missing_tbm_days(self, start, end, skip_sunday=True):
        """기간 중 TBM 기록이 없는 날짜 (일요일 제외 옵션)."""
        s, e = dt.date.fromisoformat(parse_date(start)), dt.date.fromisoformat(parse_date(end))
        have = {
            r["exec_date"]
            for r in self.db.execute(
                "SELECT DISTINCT exec_date FROM tbm WHERE exec_date BETWEEN ? AND ?",
                (s.isoformat(), e.isoformat()),
            )
        }
        out, d = [], s
        while d <= e:
            if not (skip_sunday and d.weekday() == 6) and d.isoformat() not in have:
                out.append(d.isoformat())
            d += dt.timedelta(days=1)
        return out

    def ptw_numbers(self):
        return [
            r["ptw_no"]
            for r in self.db.execute(
                "SELECT ptw_no, MAX(exec_date) m FROM tbm WHERE ptw_no<>'' GROUP BY ptw_no ORDER BY m DESC"
            )
        ]

    # ---------------------------------------------------------------- common
    def _upsert(self, table, fields, d, rid):
        keys = [k for k, _, _ in fields]
        try:
            if rid:
                self.db.execute(
                    f"UPDATE {table} SET {', '.join(k + '=?' for k in keys)} WHERE id=?",
                    [d[k] for k in keys] + [rid],
                )
            else:
                cur = self.db.execute(
                    f"INSERT INTO {table} ({', '.join(keys)}) VALUES ({', '.join('?' * len(keys))})",
                    [d[k] for k in keys],
                )
                rid = cur.lastrowid
        except sqlite3.IntegrityError:
            raise ValueError(f"이미 등록된 보고서 번호입니다: {d.get('report_no')}")
        self.db.commit()
        return rid

    def get(self, table, rid):
        r = self.db.execute(f"SELECT * FROM {table} WHERE id=?", (rid,)).fetchone()
        return dict(r) if r else None

    def delete(self, table, ids):
        self.db.executemany(f"DELETE FROM {table} WHERE id=?", [(i,) for i in ids])
        self.db.commit()

    # ---------------------------------------------------------------- import
    def import_ncr(self, rows):
        """보고서 번호가 같으면 갱신, 없으면 추가. (추가수, 갱신수, 오류목록)."""
        added = updated = 0
        errors = []
        for i, r in enumerate(rows, 1):
            try:
                no = str(r.get("report_no") or "").strip()
                old = no and self.db.execute("SELECT * FROM ncr WHERE report_no=?", (no,)).fetchone()
                if old:
                    merged = dict(old)
                    merged.update({k: v for k, v in r.items() if v not in (None, "")})
                    self.save_ncr(merged, old["id"])
                    updated += 1
                else:
                    self.save_ncr(r)
                    added += 1
            except ValueError as e:
                errors.append(f"{i}행: {e}")
        return added, updated, errors

    def import_tbm(self, rows):
        """PTW 관리번호+실시일+협력업체가 같은 건은 건너뜀. (추가수, 중복수, 오류목록)."""
        added = skipped = 0
        errors = []
        for i, r in enumerate(rows, 1):
            try:
                ptw = str(r.get("ptw_no") or "").strip()
                day = parse_date(r.get("exec_date"))
                comp = str(r.get("company") or "").strip() or self.setting("company")
                if ptw and self.db.execute(
                    "SELECT 1 FROM tbm WHERE ptw_no=? AND exec_date=? AND company=?", (ptw, day, comp)
                ).fetchone():
                    skipped += 1
                    continue
                self.save_tbm(r)
                added += 1
            except ValueError as e:
                errors.append(f"{i}행: {e}")
        return added, skipped, errors

    # --------------------------------------------------------------- summary
    def monthly_summary(self, year):
        ref = today()
        out = []
        for m in range(1, 13):
            ym = f"{year}-{m:02d}"
            ncr = [dict(r) for r in self.db.execute("SELECT * FROM ncr WHERE issue_date LIKE ?", (ym + "%",))]
            tbm_days = self.db.execute(
                "SELECT COUNT(DISTINCT exec_date), COUNT(*) FROM tbm WHERE exec_date LIKE ?", (ym + "%",)
            ).fetchone()
            rate = on_time_rate(ncr, ref)
            out.append({
                "month": ym,
                "ncr": len(ncr),
                "done": sum(1 for r in ncr if r["done_date"]),
                "open": sum(1 for r in ncr if not r["done_date"]),
                "overdue": sum(1 for r in ncr if ncr_status(r, ref) == "기한초과"),
                "rate": "" if rate is None else f"{rate:.2f}%",
                "tbm_days": tbm_days[0],
                "tbm_count": tbm_days[1],
            })
        return out


# ------------------------------------------------------------------ Excel I/O
NCR_ALIASES = {
    "보고서번호": "report_no", "발행부서": "dept", "출처": "source", "작성자": "author",
    "발행일": "issue_date", "조치요구일": "due_date", "조치완료": "done_date",
    "조치완료일": "done_date", "수신처": "receiver", "발행구분": "issue_type",
    "제목": "title", "결과확인": "result", "위치/작업": "location", "위치": "location",
    "조치내용": "action", "비고": "memo",
}
TBM_ALIASES = {
    "현장코드": "site_code", "ptw관리번호": "ptw_no", "ptw번호": "ptw_no",
    "협력업체명": "company", "협력업체": "company", "등록일": "reg_date", "실시일": "exec_date",
    "실시구분": "kind", "구분": "kind", "작업내용": "work", "작업반장": "leader",
    "참석인원": "attendees", "위험요인": "hazards", "안전대책": "measures", "비고": "memo",
}


def _norm(h):
    return re.sub(r"\s+", "", str(h or "")).lower()


def export_xlsx(path, title, fields, rows, extra=()):
    """rows(dict 목록)를 HPMS 화면과 같은 열 순서로 엑셀 저장. extra=[(키, 머리글)]."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    cols = [("_no", "No", 6)] + [(k, h, w) for k, h, w in fields] + [(k, h, 12) for k, h in extra]
    ws.cell(1, 1, title).font = Font(bold=True, size=14)
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, (_, h, w) in enumerate(cols, 1):
        cell = ws.cell(3, c, h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE6F1")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = max(6, w // 7)
    red = Font(color="C00000")
    for i, r in enumerate(rows, 1):
        for c, (k, _, _) in enumerate(cols, 1):
            cell = ws.cell(3 + i, c, i if k == "_no" else r.get(k, ""))
            cell.border = border
            if r.get("status") == "기한초과":
                cell.font = red
    ws.freeze_panes = "A4"
    if rows:
        ws.auto_filter.ref = f"A3:{ws.cell(3 + len(rows), len(cols)).coordinate}"
    wb.save(path)


def import_xlsx(path, aliases):
    """HPMS 에서 내려받은 엑셀 등을 읽어 dict 목록 반환 (머리글 이름으로 열 자동 매칭)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        rows = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()  # read_only 모드는 닫지 않으면 Windows 에서 파일이 잠김
    for hi, row in enumerate(rows[:20]):
        mapping = {i: aliases[_norm(v)] for i, v in enumerate(row) if _norm(v) in aliases}
        if len(mapping) >= 3:
            break
    else:
        raise ValueError("머리글(보고서 번호/제목 또는 PTW 관리번호/실시일 등)을 찾지 못했습니다.")
    out = []
    for row in rows[hi + 1:]:
        d = {k: row[i] for i, k in mapping.items() if i < len(row)}
        if any(v not in (None, "") for v in d.values()):
            out.append(d)
    return out
